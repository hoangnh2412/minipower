#!/usr/bin/env python3
"""XLSX báo giá nội bộ đã chốt -> HTML báo giá gửi khách hàng (in PDF thủ công từ trình duyệt).

Tái dùng đọc/tính của from_xlsx.py — không viết lại công thức. Chỉ bỏ phần giải thích
nội bộ (công thức catalog, provenance ĐO THẬT/GIẢ ĐỊNH, ROM, rủi ro, nghiệm thu…),
giữ lại Tổng quan (giá + 3 cấu phần) và Chi tiết công việc (theo nhóm, không drill-down
từng dòng mục/nguồn — trừ khi chọn --detail lines).

Không hardcode tên khách hàng/dự án: đọc từ `00-THAMSO` (qua read_params), hoặc ghi đè
bằng --khach-hang/--du-an.
"""
import argparse
import html
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

import from_xlsx as fx
import params as P

e = html.escape
vn = fx.vn
pct = fx.pct


def calc_chain_manmonth(s, don_gia_mm, tien):
    return f"""
  <div class="calc-chain">
    <div class="step">
      <div class="step-label">Khối lượng công việc</div>
      <div class="step-value">{vn(s['md'], 1)} ngày công</div>
    </div>
    <div class="op">÷ {vn(s['quydoi']['md_per_mm'], 2)}<br><span>ngày công/man-month</span></div>
    <div class="step">
      <div class="step-label">Quy đổi man-month</div>
      <div class="step-value">{vn(s['mm'], 2)} man-month</div>
    </div>
    <div class="op">× {vn(don_gia_mm)}<br><span>{e(tien)}/man-month</span></div>
    <div class="step step-final">
      <div class="step-label">Tổng giá trị báo giá</div>
      <div class="step-value">{vn(s['gia_ban'])} {e(tien)}</div>
    </div>
  </div>
  <p class="note">Khối lượng công việc (ngày công) được tính từ chi tiết công việc ở mục 2,
  quy đổi sang man-month theo năng lực làm việc thực tế ({vn(s['quydoi']['md_per_mm'], 2)}
  ngày công/man-month, đã trừ ngày lễ/phép và hệ số sử dụng thời gian), rồi nhân với đơn giá
  bán bình quân theo man-month để ra tổng giá trị báo giá.<br>
  * Đơn giá bán bình quân trên toàn cơ cấu vai trò tham gia dự án — không phải một đơn giá
  áp chung cho mọi đầu việc.</p>
"""


def calc_chain_manday(s, rate_md, tien):
    return f"""
  <div class="calc-chain">
    <div class="step">
      <div class="step-label">Khối lượng công việc</div>
      <div class="step-value">{vn(s['md'], 1)} ngày công</div>
    </div>
    <div class="op">× {vn(rate_md)}<br><span>{e(tien)}/ngày công *</span></div>
    <div class="step step-final">
      <div class="step-label">Tổng giá trị báo giá</div>
      <div class="step-value">{vn(s['gia_ban'])} {e(tien)}</div>
    </div>
  </div>
  <p class="note">Khối lượng công việc (ngày công) được tính từ chi tiết công việc ở mục 2,
  nhân thẳng với đơn giá bán bình quân theo ngày công để ra tổng giá trị báo giá.<br>
  * Đơn giá bán bình quân trên toàn cơ cấu vai trò tham gia dự án, quy đổi theo ngày công
  — không phải một đơn giá áp chung cho mọi đầu việc.</p>
"""


def cau_phan_table(s, mh_per_md):
    khung_md = s["khung"] / mh_per_md
    nguon_md = s["nguon"] / mh_per_md
    ngoai_le_md = s["ngoai_le"] / mh_per_md
    tong_md = khung_md + nguon_md + ngoai_le_md
    rows = [
        ("Phí khung — hạng mục chung, tính một lần cho toàn dự án", khung_md),
        ("Gói theo nguồn — khối lượng công việc đã xác định rõ phạm vi, cam kết giá cố định", nguon_md),
        ("Ngoại lệ — phạm vi chưa khảo sát xong hoặc chưa chốt danh mục chi tiết, tính riêng khi chốt", ngoai_le_md),
    ]
    trs = "".join(
        f"<tr><td>{e(label)}</td><td class=n>{vn(md, 1)}</td>"
        f"<td class=n>{pct(md / tong_md if tong_md else 0)}</td></tr>"
        for label, md in rows
    )
    return f"""
  <table>
    <thead><tr><th>Cấu phần</th><th class="n">Ngày công (man-day)</th><th class="n">Tỷ trọng</th></tr></thead>
    <tbody>
      {trs}
      <tr class="tot"><td>Tổng cộng</td><td class="n">{vn(tong_md, 1)}</td><td class="n">100%</td></tr>
    </tbody>
  </table>
  <p class="note">Phần "Ngoại lệ" không nằm trong cam kết giá cố định của báo giá này — sẽ
  chốt lại sau khi hoàn tất khảo sát phần nguồn dữ liệu liên quan.</p>
""" if ngoai_le_md else f"""
  <table>
    <thead><tr><th>Cấu phần</th><th class="n">Ngày công (man-day)</th><th class="n">Tỷ trọng</th></tr></thead>
    <tbody>
      {trs}
      <tr class="tot"><td>Tổng cộng</td><td class="n">{vn(tong_md, 1)}</td><td class="n">100%</td></tr>
    </tbody>
  </table>
"""


def work_detail_groups(nen, group2, g3, g4, mh_per_md):
    rows = [("Hạng mục chung — quy ước thiết kế dùng chung, phương pháp luận, khung bảo vệ dữ liệu "
             "cá nhân, họp/báo cáo định kỳ", sum(n["mh"] for n in nen) / mh_per_md)]
    rows += [(label, mh / mh_per_md) for label, _, mh in group2]
    if g3:
        rows.append((P.GROUP_BAOCAO_LABEL, sum(l["mh"] for l in g3) / mh_per_md))
    if g4:
        rows.append((P.GROUP_HOTRO_LABEL, sum(l["mh"] for l in g4) / mh_per_md))
    tong = sum(md for _, md in rows) or 1
    trs = []
    for label, md in rows:
        cls = ' class="zero"' if md == 0 else ""
        trs.append(
            f"<tr{cls}><td>{e(label)}</td><td class=n>{vn(md, 1)}</td>"
            f"<td class=n>{pct(md / tong)}</td></tr>"
        )
    return "\n      ".join(trs), tong


def work_detail_lines(nen, group2, g3, g4, mh_per_md):
    rows = [("Hạng mục chung", [("Hạng mục chung", sum(n["mh"] for n in nen) / mh_per_md)])]
    for label, g, _ in group2:
        rows.append((label, [(str(l["hang_muc"]), l["mh"] / mh_per_md) for l in g]))
    if g3:
        rows.append((P.GROUP_BAOCAO_LABEL, [(str(l["hang_muc"]), l["mh"] / mh_per_md) for l in g3]))
    if g4:
        rows.append((P.GROUP_HOTRO_LABEL, [(str(l["hang_muc"]), l["mh"] / mh_per_md) for l in g4]))
    tong = sum(md for _, items in rows for _, md in items) or 1
    blocks = []
    for label, items in rows:
        sub = sum(md for _, md in items)
        trs = "\n      ".join(
            f"<tr><td>{e(it)}</td><td class=n>{vn(md, 1)}</td><td class=n>{pct(md / tong)}</td></tr>"
            for it, md in items
        )
        blocks.append(
            f'<tr class="grp"><td colspan=3>{e(label)}</td></tr>\n      {trs}\n      '
            f'<tr class="sub"><td>Σ {e(label)}</td><td class=n>{vn(sub, 1)}</td>'
            f'<td class=n>{pct(sub / tong)}</td></tr>'
        )
    return "\n      ".join(blocks), tong


CSS = """
:root{
  --ink:#1a1a1a; --mut:#5a5a5a; --line:#d8d8d8;
  --head:#1f3864; --accent:#1f3864; --zebra:#f6f8fb; --hi:#e8f0fe; --bg:#fff;
}
@media(prefers-color-scheme:dark){:root:not([data-theme=light]){
  --ink:#e8e8e8; --mut:#9aa0a6; --line:#30343a;
  --head:#4a76c4; --accent:#9bb8e8; --zebra:#1a1d22; --hi:#1b2a44; --bg:#14161a;
}}
:root[data-theme=dark]{
  --ink:#e8e8e8; --mut:#9aa0a6; --line:#30343a;
  --head:#4a76c4; --accent:#9bb8e8; --zebra:#1a1d22; --hi:#1b2a44; --bg:#14161a;
}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);
  font:14px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,Arial,sans-serif}
.page{max-width:900px;margin:0 auto;padding:48px 56px 64px}
.brand{font-size:12px;letter-spacing:.08em;text-transform:uppercase;color:var(--mut);margin-bottom:6px}
h1{font-size:24px;margin:0 0 4px;color:var(--accent);letter-spacing:-.2px}
.meta{width:100%;border-collapse:collapse;margin-bottom:28px;font-size:13px}
.meta td{padding:4px 0;vertical-align:top}
.meta td.k{color:var(--mut);width:140px}
h2{font-size:16px;margin:34px 0 14px;padding-bottom:7px;border-bottom:2px solid var(--line);color:var(--accent)}
.price-box{display:flex;justify-content:space-between;align-items:center;flex-wrap:wrap;gap:12px;
  background:var(--hi);border:1px solid var(--line);border-radius:8px;padding:20px 26px;margin:18px 0 22px}
.price-box .label{font-size:13px;color:var(--mut);text-transform:uppercase;letter-spacing:.06em}
.price-box .value{font-size:28px;font-weight:700;color:var(--accent);margin-top:4px}
.calc-chain{display:flex;align-items:center;justify-content:center;flex-wrap:wrap;gap:10px 14px;
  margin:20px 0 10px;padding:18px 14px;background:var(--zebra);border:1px solid var(--line);border-radius:8px}
.calc-chain .step{text-align:center;min-width:140px}
.calc-chain .step-label{font-size:11px;color:var(--mut);text-transform:uppercase;letter-spacing:.05em;margin-bottom:4px}
.calc-chain .step-value{font-size:16px;font-weight:700;color:var(--ink)}
.calc-chain .step-final .step-value{color:var(--accent);font-size:18px}
.calc-chain .step-final{min-width:170px}
.calc-chain .op{font-size:13px;font-weight:600;color:var(--accent);text-align:center;line-height:1.3;white-space:nowrap;padding:0 2px}
.calc-chain .op span{display:block;font-size:10.5px;font-weight:400;color:var(--mut);white-space:normal}
table{border-collapse:collapse;width:100%;font-size:13.5px;margin:8px 0 6px}
th,td{border:1px solid var(--line);padding:8px 12px;text-align:left;vertical-align:top}
th{background:var(--head);color:#fff;font-weight:600}
td.n,th.n{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
tbody tr:nth-child(even){background:var(--zebra)}
tr.tot td{font-weight:700;background:var(--hi)}
tr.zero td{color:var(--mut);font-style:italic}
tr.grp td{font-weight:700;background:var(--zebra)}
tr.sub td{font-weight:600;background:var(--hi)}
.note{color:var(--mut);font-size:12.5px;margin:6px 0 0}
.footer{margin-top:48px;padding-top:16px;border-top:1px solid var(--line);color:var(--mut);
  font-size:11.5px;display:flex;justify-content:space-between}
@media print{body{background:#fff}.page{padding:0;max-width:none}h2{break-after:avoid}
  tr{break-inside:avoid}.price-box{break-inside:avoid}}
"""


def render(du_an, khach_hang, ngay_baogia, hieu_luc_ngay, s, don_gia_mm, rate_md, unit, tien,
           chi_tiet_html, chi_tiet_tong, mh_per_md):
    chain = (
        calc_chain_manmonth(s, don_gia_mm, tien) if unit == "manmonth"
        else calc_chain_manday(s, rate_md, tien)
    )
    return f"""<!doctype html>
<html lang="vi">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Báo giá — {e(du_an)}</title>
<style>{CSS}</style>
</head>
<body>
<div class="page">

  <div class="brand">Báo giá dịch vụ tư vấn</div>
  <h1>{e(du_an)}</h1>

  <table class="meta">
    <tr><td class="k">Khách hàng</td><td>{e(khach_hang)}</td></tr>
    <tr><td class="k">Ngày báo giá</td><td>{e(ngay_baogia)}</td></tr>
    <tr><td class="k">Hiệu lực báo giá</td><td>{hieu_luc_ngay} ngày kể từ ngày phát hành</td></tr>
  </table>

  <h2>1 · Tổng quan</h2>

  <div class="price-box">
    <div>
      <div class="label">Tổng giá trị báo giá</div>
      <div class="value">{vn(s['gia_ban'])} {e(tien)}</div>
    </div>
    <div style="text-align:right">
      <div class="label">Khối lượng quy đổi</div>
      <div style="font-size:18px;font-weight:600;margin-top:4px">{vn(s['md'], 1)} ngày công ({vn(s['mm'], 2)} man-month)</div>
    </div>
  </div>
  {chain}
  {cau_phan_table(s, mh_per_md)}

  <h2>2 · Chi tiết công việc</h2>
  <p class="note">Khối lượng công việc theo từng hạng mục — cơ sở để xác định tổng giá trị báo giá ở mục 1.</p>

  <table>
    <thead><tr><th>Hạng mục</th><th class="n">Ngày công (man-day)</th><th class="n">Tỷ trọng</th></tr></thead>
    <tbody>
      {chi_tiet_html}
      <tr class="tot"><td>Tổng cộng</td><td class="n">{vn(chi_tiet_tong, 1)}</td><td class="n">100%</td></tr>
    </tbody>
  </table>

  <div class="footer">
    <div>Báo giá — {e(du_an)}</div>
    <div>Trang 1/1</div>
  </div>

</div>
</body>
</html>
"""


def main():
    ap = argparse.ArgumentParser(description="XLSX báo giá nội bộ đã chốt -> HTML báo giá gửi khách hàng")
    ap.add_argument("--input", required=True, help="workbook .xlsx đã chốt")
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--basename", help="tên file không đuôi; mặc định theo tên xlsx")
    ap.add_argument("--unit", choices=["manmonth", "manday"], default="manmonth",
                     help="đơn giá hiển thị theo man-month (mặc định) hay man-day")
    ap.add_argument("--detail", choices=["groups", "lines"], default="groups",
                     help="chi tiết công việc theo nhóm (mặc định, không drill-down) hay từng dòng mục")
    ap.add_argument("--khach-hang", help="ghi đè tên khách hàng (mặc định đọc từ 00-THAMSO)")
    ap.add_argument("--du-an", help="ghi đè tên dự án (mặc định đọc từ 00-THAMSO)")
    ap.add_argument("--hieu-luc-ngay", type=int, default=30, help="số ngày hiệu lực báo giá (mặc định 30)")
    ap.add_argument("--overwrite", action="store_true")
    a = ap.parse_args()

    src = Path(a.input)
    if not src.exists():
        sys.exit(f"không thấy file: {src}")

    wb = load_workbook(src, data_only=False)
    need = {"00-THAMSO", "01-CATALOG", "02-ADDER", "03-DONGMUC", "04-NEN", "05-TONGHOP"}
    missing = need - set(wb.sheetnames)
    if missing:
        sys.exit(f"thiếu sheet: {', '.join(sorted(missing))}")
    cached_wb = load_workbook(src, data_only=True)

    p_file, meta = fx.read_params(wb, cached_wb)
    p = P.resolve(p_file, {})
    cat = fx.read_catalog(wb)
    adder_mh = fx.read_adders(wb)
    lines, _warn = fx.read_lines(wb, cat, adder_mh, cached_wb)
    nen = fx.read_nen(wb, P.num(p.get("so_thang"), 1) or 1)
    s = fx.rollup(lines, nen, p)

    if not s.get("gia_ban"):
        sys.exit(
            "chưa tính được Giá bán (thiếu don_gia_mm hoặc luong_thang ở `00-THAMSO`) "
            "— không thể dựng báo giá khách hàng"
        )

    khach_hang = a.khach_hang or meta.get("Khách hàng") or ""
    du_an = a.du_an or meta.get("Dự án") or ""
    if not khach_hang:
        sys.exit("thiếu tên khách hàng — truyền --khach-hang hoặc điền `Khách hàng` ở 00-THAMSO")
    if not du_an:
        sys.exit("thiếu tên dự án — truyền --du-an hoặc điền `Dự án` ở 00-THAMSO")
    ngay_baogia = str(meta.get("Ngày") or date.today().isoformat())

    tien = p.get("tien_te") or "VND"
    don_gia_mm = P.num(p.get("don_gia_mm"))
    rate_md = s.get("rate_md")

    check = s["mm"] * don_gia_mm if a.unit == "manmonth" else s["md"] * rate_md
    if abs(check - s["gia_ban"]) >= 1:
        sys.exit(
            f"lệch công thức: {a.unit} × đơn giá = {check:,.2f} ≠ Giá bán = {s['gia_ban']:,.2f} "
            "— dừng, không xuất báo giá sai"
        )

    live = [l for l in lines if not l["out_of_scope"]]
    group2, (g3, _sub3), (g4, _sub4) = fx.build_section2(live)
    mh_per_md = s["mh_per_md"]

    if a.detail == "lines":
        chi_tiet_html, chi_tiet_tong = work_detail_lines(nen, group2, g3, g4, mh_per_md)
    else:
        chi_tiet_html, chi_tiet_tong = work_detail_groups(nen, group2, g3, g4, mh_per_md)

    out_html = render(
        du_an, khach_hang, ngay_baogia, a.hieu_luc_ngay, s, don_gia_mm, rate_md, a.unit, tien,
        chi_tiet_html, chi_tiet_tong, mh_per_md,
    )

    out = Path(a.out_dir)
    out.mkdir(parents=True, exist_ok=True)
    base = a.basename or f"{src.stem}-khach-hang"
    f_html = out / f"{base}.html"
    if f_html.exists() and not a.overwrite:
        sys.exit(f"đã tồn tại: {f_html} — thêm --overwrite")
    f_html.write_text(out_html, "utf-8")
    print(f"html khách hàng → {f_html}")


if __name__ == "__main__":
    main()
